"""openpyxl을 이용한 Excel 출력"""

from __future__ import annotations

import io
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.extractor import HEADER

DOMAIN_HEADER = [
    "도메인그룹",
    "도메인명",
    "데이터타입",
    "길이",
    "인포타입",
    "비고",
]

_DOMAIN_COL_WIDTHS = {
    "도메인그룹": 16,
    "도메인명": 14,
    "길이": 8,
    "데이터타입": 22,
    "인포타입": 18,
    "비고": 28,
}
_DOMAIN_CENTER = {"데이터타입", "길이", "인포타입"}


_HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
_HEADER_FONT = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
_CELL_FONT = Font(name="맑은 고딕", size=10)
_THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
_CENTER = Alignment(horizontal="center", vertical="center")
_LEFT = Alignment(horizontal="left", vertical="center")

_COL_WIDTHS = {
    "용어명": 22,
    "용어영문명": 24,
    "용어정의": 30,
    "도메인명": 14,
    "인포타입": 18,
    "데이터타입": 18,
}

_CENTER_COLS = {"도메인명", "인포타입", "데이터타입"}


def write_excel(records: list[dict[str, str]], path: str | Path | None = None) -> bytes:
    """레코드 리스트를 Excel 파일로 작성. path가 주어지면 파일 저장, 항상 bytes도 반환."""
    wb = Workbook()
    ws = wb.active
    ws.title = "용어사전"

    for ci, h in enumerate(HEADER, start=1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.border = _THIN_BORDER
        cell.alignment = _CENTER

    for ri, rec in enumerate(records, start=2):
        for ci, h in enumerate(HEADER, start=1):
            cell = ws.cell(row=ri, column=ci, value=rec.get(h, ""))
            cell.font = _CELL_FONT
            cell.border = _THIN_BORDER
            cell.alignment = _CENTER if h in _CENTER_COLS else _LEFT

    for ci, h in enumerate(HEADER, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = _COL_WIDTHS.get(h, 15)

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    if path:
        wb.save(str(path))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def domain_rows_to_sheet_dicts(rows: list[dict]) -> list[dict[str, str]]:
    """DB/API 도메인 dict 리스트를 Excel용 한글 키 레코드로 변환."""
    out: list[dict[str, str]] = []
    for r in rows:
        lg = r.get("length")
        out.append({
            "도메인그룹": str(r.get("domain_group") or ""),
            "도메인명": str(r.get("domain_name") or ""),
            "길이": "" if lg is None else str(lg),
            "데이터타입": str(r.get("data_type") or ""),
            "인포타입": str(r.get("info_type") or ""),
            "비고": str(r.get("note") or ""),
        })
    return out


def write_domains_excel(records: list[dict[str, str]], path: str | Path | None = None) -> bytes:
    """표준도메인사전 시트 Excel."""
    wb = Workbook()
    ws = wb.active
    ws.title = "표준도메인"

    for ci, h in enumerate(DOMAIN_HEADER, start=1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.border = _THIN_BORDER
        cell.alignment = _CENTER

    for ri, rec in enumerate(records, start=2):
        for ci, h in enumerate(DOMAIN_HEADER, start=1):
            cell = ws.cell(row=ri, column=ci, value=rec.get(h, ""))
            cell.font = _CELL_FONT
            cell.border = _THIN_BORDER
            cell.alignment = _CENTER if h in _DOMAIN_CENTER else _LEFT

    for ci, h in enumerate(DOMAIN_HEADER, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = _DOMAIN_COL_WIDTHS.get(h, 14)

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    if path:
        wb.save(str(path))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
